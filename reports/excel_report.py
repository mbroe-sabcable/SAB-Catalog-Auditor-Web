from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from audits.unified_record_builder import UnifiedRecordBuilder
from reports.report_storage import get_report_directory


HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Calibri", size=11, color="1F2937")
TITLE_FONT = Font(name="Calibri", size=24, bold=True, color="1F4E78")
SUBTITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1F4E78")
THIN_SIDE = Side(style="thin", color="D9D9D9")
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)


def _set_cell(cell, value, font=None, fill=None, alignment=None, border=None):
    cell.value = value
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if alignment is not None:
        cell.alignment = alignment
    if border is not None:
        cell.border = border


def _build_output_dir():
    return get_report_directory()


def _resolve_logo_path():
    assets_dir = Path(__file__).resolve().parent.parent / "assets"
    for filename in ("sab_logo.png", "SAB_logo.png"):
        candidate = assets_dir / filename
        if candidate.exists():
            return candidate
    return None


def _insert_logo(sheet):
    logo_path = _resolve_logo_path()
    if logo_path is None:
        return

    try:
        image = Image(str(logo_path))
        image.width = 200
        image.height = 70
        sheet.add_image(image, "B2")
    except Exception:
        # Keep the dashboard generation resilient if Pillow/image loading fails.
        return


def _parse_generated_datetime(summary_data):
    generated_at = summary_data.get("generated_at", "")
    if isinstance(generated_at, datetime):
        return generated_at

    if isinstance(generated_at, str) and generated_at.strip():
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
            try:
                return datetime.strptime(generated_at.strip(), fmt)
            except ValueError:
                continue

    return datetime.now()


def _extract_source_files(summary_data):
    files = summary_data.get("source_files", {}) or {}
    return {
        "German Engineering": files.get("german_engineering") or summary_data.get("german_engineering_file", ""),
        "TrackVia": files.get("trackvia") or summary_data.get("trackvia_file", ""),
        "Directus": files.get("directus") or summary_data.get("directus_file", ""),
        "US Catalog": files.get("us_catalog") or summary_data.get("us_catalog_file", ""),
    }


def _count_corrections(summary_data):
    recommended = summary_data.get("recommended_corrections", []) or []
    if recommended:
        return len(recommended)

    total = 0
    for product in summary_data.get("product_corrections", []) or []:
        total += int(product.get("total_corrections", 0) or 0)
    return total


def _compute_health_score(products_compared, missing_products, mismatches):
    if products_compared <= 0:
        return 100

    issue_count = max(0, int(missing_products)) + max(0, int(mismatches))
    issue_ratio = issue_count / max(int(products_compared), 1)
    return max(0, min(100, round(100 - (issue_ratio * 100))))


def _health_status(score):
    if score >= 100:
        return "Excellent", PatternFill(fill_type="solid", fgColor="70AD47")
    if score >= 90:
        return "Good", PatternFill(fill_type="solid", fgColor="92D050")
    if score >= 75:
        return "Fair", PatternFill(fill_type="solid", fgColor="FFD966")
    return "Needs Attention", PatternFill(fill_type="solid", fgColor="F4B084")


def _box_header(sheet, start_col, end_col, row, title):
    sheet.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
    cell = sheet.cell(row=row, column=start_col)
    _set_cell(
        cell,
        title,
        font=HEADER_FONT,
        fill=HEADER_FILL,
        alignment=Alignment(horizontal="left", vertical="center"),
        border=THIN_BORDER,
    )
    for col in range(start_col + 1, end_col + 1):
        sheet.cell(row=row, column=col).fill = HEADER_FILL
        sheet.cell(row=row, column=col).border = THIN_BORDER


def _box_rows(sheet, start_col, end_col, start_row, items):
    row = start_row
    for label, value in items:
        sheet.merge_cells(start_row=row, start_column=start_col + 1, end_row=row, end_column=end_col)
        _set_cell(
            sheet.cell(row=row, column=start_col),
            label,
            font=Font(name="Calibri", size=11, bold=True, color="1F2937"),
            alignment=Alignment(horizontal="left", vertical="center"),
            border=THIN_BORDER,
        )
        _set_cell(
            sheet.cell(row=row, column=start_col + 1),
            value,
            font=BODY_FONT,
            alignment=Alignment(horizontal="left", vertical="center"),
            border=THIN_BORDER,
        )
        for col in range(start_col + 2, end_col + 1):
            sheet.cell(row=row, column=col).border = THIN_BORDER
        row += 1


def _autosize_columns(sheet):
    for column_cells in sheet.columns:
        max_length = 0
        col_idx = column_cells[0].column
        col_letter = get_column_letter(col_idx)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        sheet.column_dimensions[col_letter].width = min(max(max_length + 2, 12), 45)


def _style_table_sheet(sheet):
    if sheet.max_row >= 1:
        for cell in sheet[1]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = THIN_BORDER

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = THIN_BORDER

    sheet.freeze_panes = "A2"
    _autosize_columns(sheet)


def _build_dashboard_sheet(workbook, summary_data):
    dashboard = workbook.create_sheet("Dashboard")
    workbook.move_sheet(dashboard, offset=-len(workbook.sheetnames) + 1)

    for col, width in {
        "A": 4,
        "B": 22,
        "C": 22,
        "D": 22,
        "E": 22,
        "F": 4,
        "G": 22,
        "H": 22,
        "I": 22,
        "J": 22,
    }.items():
        dashboard.column_dimensions[col].width = width

    _insert_logo(dashboard)

    dashboard.merge_cells("D2:J2")
    _set_cell(
        dashboard["D2"],
        "SAB North America",
        font=TITLE_FONT,
        alignment=Alignment(horizontal="left", vertical="center"),
    )

    dashboard.merge_cells("D3:J3")
    _set_cell(
        dashboard["D3"],
        "Catalog Audit Report",
        font=SUBTITLE_FONT,
        alignment=Alignment(horizontal="left", vertical="center"),
    )

    generated_dt = _parse_generated_datetime(summary_data)
    products_compared = int(summary_data.get("matching_count", 0) or 0)
    missing_products = int(summary_data.get("missing_from_directus_count", 0) or 0) + int(summary_data.get("missing_from_trackvia_count", 0) or 0)
    mismatches = int(summary_data.get("total_mismatches", 0) or 0)
    corrections_generated = _count_corrections(summary_data)
    rubicon_weight_matches = int(summary_data.get("rubicon_weight_matches", 0) or 0)
    rubicon_weight_mismatches = int(summary_data.get("rubicon_weight_mismatches", 0) or 0)
    rubicon_not_found = int(summary_data.get("rubicon_not_found", 0) or 0)

    _box_header(dashboard, 2, 5, 6, "Report Details")
    _box_rows(
        dashboard,
        2,
        5,
        7,
        [
            ("Audit Type", summary_data.get("audit_type", "")),
            ("Date Generated", generated_dt.strftime("%Y-%m-%d")),
            ("Time Generated", generated_dt.strftime("%H:%M:%S")),
            ("Application Version", summary_data.get("app_version", "1.0.0")),
        ],
    )

    _box_header(dashboard, 7, 10, 6, "Source Files")
    source_files = _extract_source_files(summary_data)
    _box_rows(
        dashboard,
        7,
        10,
        7,
        [
            ("German Engineering", source_files.get("German Engineering", "") or ""),
            ("TrackVia", source_files.get("TrackVia", "") or ""),
            ("Directus", source_files.get("Directus", "") or ""),
            ("US Catalog", source_files.get("US Catalog", "") or ""),
        ],
    )

    _box_header(dashboard, 2, 5, 13, "Summary")
    _box_rows(
        dashboard,
        2,
        5,
        14,
        [
            ("Products Compared", products_compared),
            ("Missing Products", missing_products),
            ("Specification Mismatches", mismatches),
            ("Corrections Generated", corrections_generated),
            ("Rubicon Weight Matches", rubicon_weight_matches),
            ("Rubicon Weight Mismatches", rubicon_weight_mismatches),
            ("Not Found in Rubicon", rubicon_not_found),
        ],
    )

    _box_header(dashboard, 7, 10, 13, "Audit Health")
    health_score = _compute_health_score(products_compared, missing_products, mismatches)
    health_label, health_fill = _health_status(health_score)
    _box_rows(
        dashboard,
        7,
        10,
        14,
        [
            ("Health Score", f"{health_score}%"),
            ("Status", health_label),
        ],
    )
    for cell_ref in ("H14", "H15"):
        dashboard[cell_ref].fill = health_fill
        dashboard[cell_ref].font = Font(name="Calibri", size=11, bold=True, color="1F2937")

    _box_header(dashboard, 2, 5, 22, "Worksheet Index")
    dashboard.merge_cells("C23:E23")
    _set_cell(
        dashboard["B23"],
        "Dashboard",
        font=BODY_FONT,
        alignment=Alignment(horizontal="left", vertical="center"),
        border=THIN_BORDER,
    )
    _set_cell(
        dashboard["C23"],
        "Go to sheet",
        font=Font(name="Calibri", size=11, color="0563C1", underline="single"),
        alignment=Alignment(horizontal="left", vertical="center"),
        border=THIN_BORDER,
    )
    dashboard["C23"].hyperlink = "#'Dashboard'!A1"
    for col in range(4, 6):
        dashboard.cell(row=23, column=col).border = THIN_BORDER

    dashboard.freeze_panes = "A7"
    return dashboard


def _populate_dashboard_index(dashboard, workbook):
    start_row = 24
    index_row = start_row
    for sheet_name in workbook.sheetnames:
        if sheet_name == "Dashboard":
            continue

        dashboard.merge_cells(start_row=index_row, start_column=3, end_row=index_row, end_column=5)
        _set_cell(
            dashboard.cell(row=index_row, column=2),
            sheet_name,
            font=BODY_FONT,
            alignment=Alignment(horizontal="left", vertical="center"),
            border=THIN_BORDER,
        )
        link_cell = dashboard.cell(row=index_row, column=3)
        _set_cell(
            link_cell,
            "Go to sheet",
            font=Font(name="Calibri", size=11, color="0563C1", underline="single"),
            alignment=Alignment(horizontal="left", vertical="center"),
            border=THIN_BORDER,
        )
        link_cell.hyperlink = f"#'{sheet_name}'!A1"

        for col in range(4, 6):
            dashboard.cell(row=index_row, column=col).border = THIN_BORDER

        index_row += 1


def _append_product_corrections_sheet(workbook, product_corrections):
    product_sheet = workbook.create_sheet("Product Corrections")
    product_sheet.append(["SKU", "Total Corrections", "Category", "System", "Field", "Current Value", "Correct Value", "Reason"])

    for product in product_corrections:
        sku = UnifiedRecordBuilder.normalize_part_number(product.get("sku", ""))
        total = product.get("total_corrections", 0)
        categories = [
            ("Product Identity corrections", product.get("product_identity_corrections", [])),
            ("Engineering corrections", product.get("engineering_corrections", [])),
            ("Website QA corrections", product.get("website_qa_corrections", [])),
            ("Directus QA corrections", product.get("directus_qa_corrections", [])),
        ]

        if total == 0:
            product_sheet.append([sku, 0, "No Corrections", "", "", "", "", "No Corrections"])
            continue

        first_row_for_sku = True
        for category_name, corrections in categories:
            if not corrections:
                product_sheet.append([
                    sku if first_row_for_sku else "",
                    total if first_row_for_sku else "",
                    category_name,
                    "",
                    "",
                    "",
                    "",
                    "No Corrections",
                ])
                first_row_for_sku = False
                continue

            for correction in corrections:
                product_sheet.append(
                    [
                        sku if first_row_for_sku else "",
                        total if first_row_for_sku else "",
                        category_name,
                        correction.get("system", ""),
                        correction.get("field", ""),
                        UnifiedRecordBuilder.normalize_part_number(correction.get("current_value", "")),
                        UnifiedRecordBuilder.normalize_part_number(correction.get("correct_value", "")),
                        correction.get("reason", ""),
                    ]
                )
                first_row_for_sku = False


def write_audit_report(summary_data):
    normalize_part_number = UnifiedRecordBuilder.normalize_part_number

    output_dir = _build_output_dir()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    report_path = output_dir / f"Audit_{timestamp}.xlsx"

    workbook = Workbook()
    workbook.remove(workbook.active)

    dashboard_sheet = _build_dashboard_sheet(workbook, summary_data)

    missing_products_sheet = workbook.create_sheet("Missing Products")
    missing_products_sheet.append(["SKU", "Missing From"])
    for sku in summary_data.get("missing_from_directus", []):
        missing_products_sheet.append([normalize_part_number(sku), "Directus"])
    for sku in summary_data.get("missing_from_trackvia", []):
        missing_products_sheet.append([normalize_part_number(sku), "TrackVia"])

    mismatches_sheet = workbook.create_sheet("Specification Mismatches")
    mismatches_sheet.append(["SKU", "Field", "TrackVia Value", "Directus Value"])
    for mismatch in summary_data.get("mismatches", []):
        mismatches_sheet.append(
            [
                normalize_part_number(mismatch.get("sku", "")),
                mismatch.get("field_display", ""),
                normalize_part_number(mismatch.get("trackvia_value", "")),
                normalize_part_number(mismatch.get("directus_value", "")),
            ]
        )

    mapping_sheet = workbook.create_sheet("German-US Mapping")
    mapping_sheet.append(["German Part Number", "US Part Number"])
    for german_number, us_number in summary_data.get("part_number_mapping", {}).items():
        mapping_sheet.append([
            normalize_part_number(german_number, preserve_leading_zeros=True),
            normalize_part_number(us_number),
        ])

    four_way_sheet = workbook.create_sheet("Four-Way Comparison")
    four_way_sheet.append(["SKU", "Field", "German Engineering", "TrackVia", "Directus", "US Catalog", "Status", "Reason"])
    four_way_comparisons = summary_data.get("four_way_comparisons", [])
    awg_comparison = next((comparison for comparison in four_way_comparisons if comparison.get("field") == "AWG"), None)
    if awg_comparison is not None:
        print(awg_comparison)

    conductors_comparison = next((comparison for comparison in four_way_comparisons if comparison.get("field") == "Conductors"), None)
    if conductors_comparison is not None:
        print(conductors_comparison)

    for comparison in summary_data.get("four_way_comparisons", []):
        field_name = comparison.get("field", "")
        is_part_number_field = field_name in {"Part Number", "US Part Number", "German Part Number"}

        sku_value = comparison.get("sku", "")
        german_value = comparison.get("german", "")
        trackvia_value = comparison.get("trackvia", "")
        directus_value = comparison.get("directus", "")
        us_catalog_value = comparison.get("us_catalog", "")

        if is_part_number_field:
            sku_value = normalize_part_number(sku_value)
            german_value = normalize_part_number(german_value, preserve_leading_zeros=True)
            trackvia_value = normalize_part_number(trackvia_value, preserve_leading_zeros=True)
            if normalize_part_number(directus_value, preserve_leading_zeros=True) != "N/A":
                directus_value = normalize_part_number(directus_value)
            if normalize_part_number(us_catalog_value, preserve_leading_zeros=True) != "N/A":
                us_catalog_value = normalize_part_number(us_catalog_value)

        four_way_sheet.append(
            [
                sku_value,
                field_name,
                german_value,
                trackvia_value,
                directus_value,
                us_catalog_value,
                comparison.get("status", ""),
                comparison.get("reason", ""),
            ]
        )

    rubicon_sheet = workbook.create_sheet("Rubicon Weight Audit")
    rubicon_sheet.append(["US Part #", "Engineering Weight", "Rubicon Weight", "Status"])
    rubicon_weight_results = summary_data.get("rubicon_weight_results", []) or []
    rubicon_exceptions = [
        result for result in rubicon_weight_results
        if result.get("result") in {"WEIGHT MISMATCH", "NOT FOUND IN RUBICON"}
    ]

    if rubicon_exceptions:
        for result in rubicon_exceptions:
            rubicon_sheet.append(
                [
                    normalize_part_number(result.get("part_number", ""), preserve_leading_zeros=True),
                    result.get("engineering_weight", ""),
                    result.get("rubicon_weight", ""),
                    result.get("result", ""),
                ]
            )
    else:
        rubicon_sheet.append(["No Rubicon weight discrepancies found.", "", "", ""])

    if summary_data.get("product_corrections"):
        _append_product_corrections_sheet(workbook, summary_data.get("product_corrections", []))

    _populate_dashboard_index(dashboard_sheet, workbook)

    for sheet in workbook.worksheets:
        if sheet.title != "Dashboard":
            _style_table_sheet(sheet)

    _autosize_columns(dashboard_sheet)

    workbook.save(report_path)
    if not report_path.exists():
        raise RuntimeError(f"Workbook was not created: {report_path}")
    return report_path.name
