from typing import Any, Optional

from audits.outlier_detector import OutlierDetector
from audits.value_normalizer import ValueNormalizer


def _normalize_value(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _find_matching_column(df, expected_column_name: Optional[str]):
    if df is None or not expected_column_name:
        return None

    expected = _normalize_value(expected_column_name).lower()
    for column in df.columns:
        if _normalize_value(column).lower() == expected:
            return column
    return None


def _get_source_value(df, row, source_column_name, source_field_name):
    if df is None or row is None or not source_field_name:
        return ""

    column_name = _find_matching_column(df, source_column_name)
    if column_name is None:
        return ""

    value = row.get(column_name)
    return _normalize_value(value)


def _resolve_conductors_value(df, row, field_mapping, source_name):
    if df is None or row is None or not field_mapping:
        return ""

    conductor_fields = field_mapping.get(f"{source_name}_fields", [])
    if not conductor_fields:
        return ""

    values = []
    for field_name in conductor_fields:
        column_name = _find_matching_column(df, field_name)
        if column_name is None:
            continue
        value = _normalize_value(row.get(column_name))
        values.append(value)

    if len(values) == 2:
        no_ground_value, with_ground_value = values
        if no_ground_value and with_ground_value:
            return ""
        if no_ground_value:
            return no_ground_value
        if with_ground_value:
            return with_ground_value

    return ""


def _get_mapped_value_from_row_dict(df, row_dict, mapped_column_name):
    if df is None or row_dict is None or not mapped_column_name:
        return ""

    column_name = _find_matching_column(df, mapped_column_name)
    if column_name is None:
        return ""

    return _normalize_value(row_dict.get(column_name))


def build_data_quality_issues(records, field_mapper, trackvia_df=None, directus_df=None, german_df=None, us_catalog_df=None):
    issues = []

    if not records:
        return issues

    for record in records:
        sku = _normalize_value(getattr(record, "sku", ""))
        if not sku:
            continue

        for field in field_mapper.get_specification_mappings():
            if field.get("key") != "conductors":
                continue

            trackvia_row = getattr(record, "trackvia_row", None)
            directus_row = getattr(record, "directus_row", None)

            if trackvia_row is not None:
                no_ground_column = _find_matching_column(trackvia_df, "part_cond_no_ground")
                with_ground_column = _find_matching_column(trackvia_df, "part_cond_with_ground")
                no_ground_value = _normalize_value(trackvia_row.get(no_ground_column)) if no_ground_column is not None else ""
                with_ground_value = _normalize_value(trackvia_row.get(with_ground_column)) if with_ground_column is not None else ""
                if no_ground_value and with_ground_value:
                    issues.append({"sku": sku, "system": "TrackVia", "issue": "Both conductors fields are populated."})

            if directus_row is not None:
                no_ground_column = _find_matching_column(directus_df, "part_cond_no_ground")
                with_ground_column = _find_matching_column(directus_df, "part_cond_with_ground")
                no_ground_value = _normalize_value(directus_row.get(no_ground_column)) if no_ground_column is not None else ""
                with_ground_value = _normalize_value(directus_row.get(with_ground_column)) if with_ground_column is not None else ""
                if no_ground_value and with_ground_value:
                    issues.append({"sku": sku, "system": "Directus", "issue": "Both conductors fields are populated."})

    return issues


def _flatten_data_quality_issues(comparisons):
    issues = []
    for comparison in comparisons:
        for issue in comparison.get("data_quality_issues", []) or []:
            issues.append(issue)
    return issues


def _build_data_quality_issues_html(issues):
    if not issues:
        return ""

    issue_rows = "".join(
        f"<tr><td>{issue.get('sku', '')}</td><td>{issue.get('system', '')}</td><td>{issue.get('issue', '')}</td></tr>"
        for issue in issues
    )
    return (
        '<div class="card shadow-sm mt-4">'
        '<div class="card-body">'
        '<h2 class="h5 mb-3">Data Quality Issues</h2>'
        f'<div class="table-responsive"><table class="table table-striped align-middle"><thead><tr><th>SKU</th><th>System</th><th>Issue</th></tr></thead><tbody>{issue_rows}</tbody></table></div>'
        '</div></div>'
    )


def _build_outlier_analysis_html(outlier_rows):
    if not outlier_rows:
        return ""

    outlier_rows_html = "".join(
        f"<tr><td>{row.get('sku', '')}</td><td>{row.get('field', '')}</td><td>{row.get('german_value', '')}</td><td>{row.get('trackvia_value', '')}</td><td>{row.get('directus_value', '')}</td><td>{row.get('us_catalog_value', '')}</td><td>{row.get('analysis', {}).get('status', '')}</td><td>{row.get('analysis', {}).get('consensus_value', '')}</td><td>{row.get('analysis', {}).get('outlier_system', '')}</td><td>{row.get('analysis', {}).get('recommendation', '')}</td><td>{row.get('analysis', {}).get('confidence', '')}</td></tr>"
        for row in outlier_rows[:20]
    )
    return (
        '<div class="card shadow-sm mt-4">'
        '<div class="card-body">'
        '<h2 class="h5 mb-3">Outlier Analysis</h2>'
        f'<div class="table-responsive"><table class="table table-striped align-middle"><thead><tr><th>SKU</th><th>Field</th><th>German</th><th>TrackVia</th><th>Directus</th><th>US Catalog</th><th>Status</th><th>Consensus Value</th><th>Outlier System</th><th>Recommendation</th><th>Confidence</th></tr></thead><tbody>{outlier_rows_html}</tbody></table></div>'
        '</div></div>'
    )


def _patch_audit_result_hooks():
    try:
        import audits.audit_engine as audit_engine_module
        import reports.excel_report as excel_report_module
    except Exception:
        return

    if getattr(audit_engine_module, "_four_way_hooks_installed", False):
        return

    original_html_method = audit_engine_module.AuditEngine._build_part_number_comparison_html

    def patched_html_method(self):
        html = original_html_method(self)
        comparisons = build_four_way_comparisons(
            self.records,
            self.field_mapper,
            trackvia_df=self.trackvia_df,
            directus_df=self.directus_df,
            german_df=self.german_df,
            us_catalog_df=self.us_catalog_df,
        )
        issues_html = _build_data_quality_issues_html(_flatten_data_quality_issues(comparisons))
        outlier_html = _build_outlier_analysis_html(
            [comparison for comparison in comparisons if comparison.get("analysis", {}).get("status") not in {"PASS", ""}]
        )
        return html + issues_html + outlier_html

    audit_engine_module.AuditEngine._build_part_number_comparison_html = patched_html_method

    from datetime import datetime
    from pathlib import Path

    from openpyxl import Workbook

    def patched_excel_writer(summary_data):
        output_dir = Path(__file__).resolve().parent.parent / "reports" / "output"
        output_dir.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        report_path = output_dir / f"Audit_{timestamp}.xlsx"

        workbook = Workbook()
        workbook.remove(workbook.active)

        summary_sheet = workbook.create_sheet("Summary")
        summary_sheet.append(["Field", "Value"])
        summary_sheet.append(["Audit Type", summary_data.get("audit_type", "")])
        summary_sheet.append(["Date/Time", summary_data.get("generated_at", "")])
        summary_sheet.append(["Matching SKUs", summary_data.get("matching_count", 0)])
        summary_sheet.append(["Missing from Directus", summary_data.get("missing_from_directus_count", 0)])
        summary_sheet.append(["Missing from TrackVia", summary_data.get("missing_from_trackvia_count", 0)])
        summary_sheet.append(["Total Specification Mismatches", summary_data.get("total_mismatches", 0)])
        summary_sheet.append([])
        summary_sheet.append(["Tier 1 Field Summary", ""])
        for field_name, details in summary_data.get("tier1_summary", {}).items():
            summary_sheet.append(
                [
                    field_name,
                    f"comparisons={details.get('comparisons', 0)}; mismatches={details.get('mismatches', 0)}; status={details.get('status', 'PASS')}",
                ]
            )

        missing_products_sheet = workbook.create_sheet("Missing Products")
        missing_products_sheet.append(["SKU", "Missing From"])
        for sku in summary_data.get("missing_from_directus", []):
            missing_products_sheet.append([sku, "Directus"])
        for sku in summary_data.get("missing_from_trackvia", []):
            missing_products_sheet.append([sku, "TrackVia"])

        mismatches_sheet = workbook.create_sheet("Specification Mismatches")
        mismatches_sheet.append(["SKU", "Field", "TrackVia Value", "Directus Value"])
        for mismatch in summary_data.get("mismatches", []):
            mismatches_sheet.append(
                [
                    mismatch.get("sku", ""),
                    mismatch.get("field_display", ""),
                    mismatch.get("trackvia_value", ""),
                    mismatch.get("directus_value", ""),
                ]
            )

        mapping_sheet = workbook.create_sheet("German-US Mapping")
        mapping_sheet.append(["German Part Number", "US Part Number"])
        for german_number, us_number in summary_data.get("part_number_mapping", {}).items():
            mapping_sheet.append([german_number, us_number])

        corrections_sheet = workbook.create_sheet("Recommended Corrections")
        corrections_sheet.append(["SKU", "System", "Field", "Current Value", "Correct Value", "Reason"])
        for correction in summary_data.get("recommended_corrections", []):
            corrections_sheet.append(
                [
                    correction.get("sku", ""),
                    correction.get("system", ""),
                    correction.get("field", ""),
                    correction.get("current_value", ""),
                    correction.get("correct_value", ""),
                    correction.get("reason", ""),
                ]
            )

        four_way_sheet = workbook.create_sheet("Four-Way Comparison")
        four_way_sheet.append(["SKU", "Field", "German", "TrackVia", "Directus", "US Catalog", "Status"])
        for comparison in summary_data.get("four_way_comparisons", []):
            four_way_sheet.append(
                [
                    comparison.get("sku", ""),
                    comparison.get("field", ""),
                    comparison.get("german_value", ""),
                    comparison.get("trackvia_value", ""),
                    comparison.get("directus_value", ""),
                    comparison.get("us_catalog_value", ""),
                    comparison.get("status", ""),
                ]
            )

        quality_sheet = workbook.create_sheet("Data Quality Issues")
        quality_sheet.append(["SKU", "System", "Issue"])
        issues = _flatten_data_quality_issues(summary_data.get("four_way_comparisons", []))
        for issue in issues:
            quality_sheet.append([issue.get("sku", ""), issue.get("system", ""), issue.get("issue", "")])

        workbook.save(report_path)
        return report_path.name

    excel_report_module.write_audit_report = patched_excel_writer
    audit_engine_module.write_audit_report = patched_excel_writer
    audit_engine_module._four_way_hooks_installed = True


def build_four_way_comparisons(records, field_mapper, trackvia_df=None, directus_df=None, german_df=None, us_catalog_df=None):
    _patch_audit_result_hooks()
    comparisons = []

    if not records:
        return comparisons

    data_quality_issues = build_data_quality_issues(
        records,
        field_mapper,
        trackvia_df=trackvia_df,
        directus_df=directus_df,
        german_df=german_df,
        us_catalog_df=us_catalog_df,
    )

    for record in records:
        sku = _normalize_value(getattr(record, "sku", ""))
        if not sku:
            continue

        trackvia_row = getattr(record, "trackvia_row", None)
        directus_row = getattr(record, "directus_row", None)
        german_row = getattr(record, "german_row", None)
        us_catalog_row = getattr(record, "us_catalog_row", None)
        sku_data_quality_issues = [issue for issue in data_quality_issues if issue.get("sku") == sku]

        for field in field_mapper.get_specification_mappings():
            field_key = field.get("key")
            field_display = field.get("display")
            if not field_key or not field_display:
                continue

            german_value = ""
            trackvia_value = ""
            directus_value = ""
            us_catalog_value = ""

            if field.get("key") == "conductors":
                if german_row is not None:
                    german_value = _get_mapped_value_from_row_dict(german_df, german_row, field.get("german"))

                if trackvia_row is not None:
                    trackvia_value = _resolve_conductors_value(trackvia_df, trackvia_row, field, "trackvia")

                if directus_row is not None:
                    directus_value = _resolve_conductors_value(directus_df, directus_row, field, "directus")

                if us_catalog_row is not None:
                    us_catalog_value = _get_mapped_value_from_row_dict(us_catalog_df, us_catalog_row, field.get("us_catalog"))
            else:
                if german_row is not None:
                    german_value = _get_mapped_value_from_row_dict(german_df, german_row, field.get("german"))

                if trackvia_row is not None:
                    trackvia_value = _get_mapped_value_from_row_dict(trackvia_df, trackvia_row, field.get("trackvia"))

                if directus_row is not None:
                    directus_value = _get_mapped_value_from_row_dict(directus_df, directus_row, field.get("directus"))

                if us_catalog_row is not None:
                    us_catalog_value = _get_mapped_value_from_row_dict(us_catalog_df, us_catalog_row, field.get("us_catalog"))

            normalized_values = [
                ValueNormalizer.normalize(value)
                for value in [german_value, trackvia_value, directus_value, us_catalog_value]
            ]
            available_values = [value for value in normalized_values if value]
            if not available_values:
                status = "PASS"
            else:
                status = "PASS" if all(value == available_values[0] for value in available_values[1:]) else "FAIL"

            analyze_args = [
                {"system": "German", "value": german_value},
                {"system": "TrackVia", "value": trackvia_value},
                {"system": "Directus", "value": directus_value},
                {"system": "US Catalog", "value": us_catalog_value},
            ]

            analysis = OutlierDetector.analyze(analyze_args)

            comparisons.append(
                {
                    "sku": sku,
                    "field": field_display,
                    "german_value": german_value,
                    "trackvia_value": trackvia_value,
                    "directus_value": directus_value,
                    "us_catalog_value": us_catalog_value,
                    "status": status,
                    "analysis": analysis,
                    "data_quality_issues": sku_data_quality_issues,
                }
            )

    return comparisons
