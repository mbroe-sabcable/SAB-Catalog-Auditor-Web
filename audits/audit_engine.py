from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

import audits.four_way_comparison as four_way_comparison_module
from audits.four_way_comparison import FourWayComparison
from audits.part_number import build_part_number_mapping
from audits.specifications import compare_specifications
from audits.unified_record_builder import AuditRecord, UnifiedRecordBuilder
from importers.field_mapper import FieldMapper
from reports import excel_report as excel_report_module
from reports.correction_report import build_correction_recommendations


def _build_outlier_analysis_html(outlier_rows):
    if not outlier_rows:
        return ""

    rows_html = "".join(
        f"<tr><td>{row.get('sku', '')}</td><td>{row.get('field', '')}</td><td>{row.get('german_value', '')}</td><td>{row.get('trackvia_value', '')}</td><td>{row.get('directus_value', '')}</td><td>{row.get('us_catalog_value', '')}</td><td>{row.get('analysis', {}).get('status', '')}</td><td>{row.get('analysis', {}).get('source_of_truth', '')}</td><td>{'; '.join(row.get('analysis', {}).get('systems_out_of_sync', []))}</td><td>{row.get('analysis', {}).get('correct_value', '')}</td><td>{'; '.join(row.get('analysis', {}).get('recommended_updates', []))}</td></tr>"
        for row in outlier_rows[:20]
    )
    return (
        '<div class="card shadow-sm mt-4">'
        '<div class="card-body">'
        '<h2 class="h5 mb-3">Outlier Analysis</h2>'
        f'<div class="table-responsive"><table class="table table-striped align-middle"><thead><tr><th>SKU</th><th>Field</th><th>German</th><th>TrackVia</th><th>Directus</th><th>US Catalog</th><th>Status</th><th>Source of Truth</th><th>Systems Out of Sync</th><th>Correct Value</th><th>Recommended Updates</th></tr></thead><tbody>{rows_html}</tbody></table></div>'
        '</div></div>'
    )


def _append_product_corrections_sheet(workbook, product_corrections):
    product_sheet = workbook.create_sheet("Product Corrections")
    product_sheet.append(["SKU", "Total Corrections", "Category", "System", "Field", "Current Value", "Correct Value", "Reason"])

    for product in product_corrections:
        sku = product.get("sku", "")
        total = product.get("total_corrections", 0)
        categories = [
            ("Product Identity corrections", product.get("product_identity_corrections", [])),
            ("Engineering corrections", product.get("engineering_corrections", [])),
            ("Website QA corrections", product.get("website_qa_corrections", [])),
            ("Directus QA corrections", product.get("directus_qa_corrections", [])),
        ]

        if total == 0:
            product_sheet.append([sku, 0, "No Corrections", "", "", "", "", "No Corrections"])
            continue

        first_row_for_sku = True
        for category_name, corrections in categories:
            if not corrections:
                product_sheet.append([
                    sku if first_row_for_sku else "",
                    total if first_row_for_sku else "",
                    category_name,
                    "",
                    "",
                    "",
                    "",
                    "No Corrections",
                ])
                first_row_for_sku = False
                continue

            for correction in corrections:
                product_sheet.append(
                    [
                        sku if first_row_for_sku else "",
                        total if first_row_for_sku else "",
                        category_name,
                        correction.get("system", ""),
                        correction.get("field", ""),
                        correction.get("current_value", ""),
                        correction.get("correct_value", ""),
                        correction.get("reason", ""),
                    ]
                )
                first_row_for_sku = False


def _write_audit_report_with_outlier_columns(summary_data):
    output_dir = Path(__file__).resolve().parent.parent / "reports" / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    report_path = output_dir / f"Audit_{timestamp}.xlsx"

    workbook = Workbook()
    workbook.remove(workbook.active)

    summary_sheet = workbook.create_sheet("Summary")
    summary_sheet.append(["Field", "Value"])
    summary_sheet.append(["Audit Type", summary_data.get("audit_type", "")])
    summary_sheet.append(["Date/Time", summary_data.get("generated_at", "")])
    summary_sheet.append(["Matching SKUs", summary_data.get("matching_count", 0)])
    summary_sheet.append(["Missing from Directus", summary_data.get("missing_from_directus_count", 0)])
    summary_sheet.append(["Missing from TrackVia", summary_data.get("missing_from_trackvia_count", 0)])
    summary_sheet.append(["Total Specification Mismatches", summary_data.get("total_mismatches", 0)])
    summary_sheet.append([])
    summary_sheet.append(["Tier 1 Field Summary", ""])
    for field_name, details in summary_data.get("tier1_summary", {}).items():
        summary_sheet.append(
            [
                field_name,
                f"comparisons={details.get('comparisons', 0)}; mismatches={details.get('mismatches', 0)}; status={details.get('status', 'PASS')}",
            ]
        )

    missing_products_sheet = workbook.create_sheet("Missing Products")
    missing_products_sheet.append(["SKU", "Missing From"])
    for sku in summary_data.get("missing_from_directus", []):
        missing_products_sheet.append([sku, "Directus"])
    for sku in summary_data.get("missing_from_trackvia", []):
        missing_products_sheet.append([sku, "TrackVia"])

    mismatches_sheet = workbook.create_sheet("Specification Mismatches")
    mismatches_sheet.append(["SKU", "Field", "TrackVia Value", "Directus Value"])
    for mismatch in summary_data.get("mismatches", []):
        mismatches_sheet.append(
            [
                mismatch.get("sku", ""),
                mismatch.get("field_display", ""),
                mismatch.get("trackvia_value", ""),
                mismatch.get("directus_value", ""),
            ]
        )

    mapping_sheet = workbook.create_sheet("German-US Mapping")
    mapping_sheet.append(["German Part Number", "US Part Number"])
    for german_number, us_number in summary_data.get("part_number_mapping", {}).items():
        mapping_sheet.append([german_number, us_number])

    corrections_sheet = workbook.create_sheet("Recommended Corrections")
    corrections_sheet.append(["SKU", "System", "Field", "Current Value", "Correct Value", "Reason"])
    for correction in summary_data.get("recommended_corrections", []):
        corrections_sheet.append(
            [
                correction.get("sku", ""),
                correction.get("system", ""),
                correction.get("field", ""),
                correction.get("current_value", ""),
                correction.get("correct_value", ""),
                correction.get("reason", ""),
            ]
        )

    outlier_sheet = workbook.create_sheet("Outlier Analysis")
    outlier_sheet.append(["SKU", "Field", "German", "TrackVia", "Directus", "US Catalog", "Status", "Source of Truth", "Systems Out of Sync", "Correct Value", "Recommended Updates"])
    for comparison in summary_data.get("four_way_comparisons", []):
        analysis = comparison.get("analysis", {})
        if analysis.get("status") in {"PASS", ""}:
            continue
        outlier_sheet.append(
            [
                comparison.get("sku", ""),
                comparison.get("field", ""),
                comparison.get("german_value", ""),
                comparison.get("trackvia_value", ""),
                comparison.get("directus_value", ""),
                comparison.get("us_catalog_value", ""),
                analysis.get("status", ""),
                analysis.get("source_of_truth", ""),
                "; ".join(analysis.get("systems_out_of_sync", [])),
                analysis.get("correct_value", ""),
                "; ".join(analysis.get("recommended_updates", [])),
            ]
        )

    _append_product_corrections_sheet(workbook, summary_data.get("product_corrections", []))

    workbook.save(report_path)
    return report_path.name


if not getattr(four_way_comparison_module, "_outlier_reporting_patched", False):
    four_way_comparison_module._build_outlier_analysis_html = _build_outlier_analysis_html
    four_way_comparison_module._outlier_reporting_patched = True


class AuditEngine:
    def __init__(self, trackvia_df=None, directus_df=None, german_df=None, us_catalog_df=None, audit_type="Full Product Family Audit"):
        self.trackvia_df = trackvia_df
        self.directus_df = directus_df
        self.german_df = german_df
        self.us_catalog_df = us_catalog_df
        self.audit_type = audit_type
        self.field_mapper = FieldMapper()
        self.records = UnifiedRecordBuilder.build(trackvia_df, directus_df, german_df, us_catalog_df)
        self.four_way_comparisons = []

    def _write_debug_file(self):
        debug_path = Path(__file__).resolve().parent.parent / "reports" / "debug.txt"
        debug_path.parent.mkdir(parents=True, exist_ok=True)

        german_columns = self.german_df.columns.tolist() if self.german_df is not None else []
        us_catalog_columns = self.us_catalog_df.columns.tolist() if self.us_catalog_df is not None else []

        first_record = self.records[0] if self.records else None
        first_record_sku = getattr(first_record, "sku", None) if first_record is not None else None
        first_record_german_part = getattr(first_record, "german_part", None) if first_record is not None else None
        first_record_trackvia_row = getattr(first_record, "trackvia_row", None) if first_record is not None else None
        first_record_german_row = getattr(first_record, "german_row", None) if first_record is not None else None
        first_record_us_catalog_row = getattr(first_record, "us_catalog_row", None) if first_record is not None else None

        german_lookup_key = first_record_german_part if first_record_german_part else first_record_sku

        us_lookup_key = first_record_sku
        if isinstance(first_record_trackvia_row, dict) and self.trackvia_df is not None:
            trackvia_column = next((column for column in self.trackvia_df.columns if str(column).lower() == "sku"), None)
            if trackvia_column is not None:
                raw_us_lookup_key = first_record_trackvia_row.get(trackvia_column)
                if raw_us_lookup_key is not None:
                    us_lookup_key_text = str(raw_us_lookup_key).strip()
                    us_lookup_key = us_lookup_key_text if us_lookup_key_text else None

        awg_mapping = self.field_mapper.get_mapping("awg") or {}
        german_awg_field = awg_mapping.get("german")
        us_catalog_awg_field = awg_mapping.get("us_catalog")

        german_awg_key_exists = bool(
            isinstance(first_record_german_row, dict)
            and german_awg_field
            and german_awg_field in first_record_german_row
        )
        us_catalog_awg_key_exists = bool(
            isinstance(first_record_us_catalog_row, dict)
            and us_catalog_awg_field
            and us_catalog_awg_field in first_record_us_catalog_row
        )

        german_awg_value = (
            first_record_german_row.get(german_awg_field, "")
            if german_awg_key_exists
            else ""
        )
        us_catalog_awg_value = (
            first_record_us_catalog_row.get(us_catalog_awg_field, "")
            if us_catalog_awg_key_exists
            else ""
        )

        lines = [
            f"German DataFrame columns: {german_columns}",
            f"US Catalog DataFrame columns: {us_catalog_columns}",
            f"First AuditRecord: {first_record}",
            f"german_row populated: {first_record_german_row is not None}",
            f"us_catalog_row populated: {first_record_us_catalog_row is not None}",
            f"German lookup key: {german_lookup_key}",
            f"US Catalog lookup key: {us_lookup_key}",
            "AWG lookup information:",
            f"  German field name: {german_awg_field}",
            f"  German key exists in row: {german_awg_key_exists}",
            f"  German value returned: {german_awg_value}",
            f"  US Catalog field name: {us_catalog_awg_field}",
            f"  US Catalog key exists in row: {us_catalog_awg_key_exists}",
            f"  US Catalog value returned: {us_catalog_awg_value}",
        ]

        debug_path.write_text("\n".join(lines) + "\n", encoding="utf-8")

    def _normalize_column_name(self, column_name):
        return "".join(ch.lower() for ch in str(column_name) if ch.isalnum())

    def _unique_values(self, df, column_name):
        values = []
        seen = set()
        for value in df[column_name].dropna():
            text = str(value).strip()
            if text and text not in seen:
                seen.add(text)
                values.append(text)
        return values

    def _get_mapped_column(self, df, source_name):
        if df is None:
            return None

        expected_column = self.field_mapper.get_part_number_mapping(source_name)
        for column in df.columns:
            if self._normalize_column_name(column) == self._normalize_column_name(expected_column):
                return column
        return None

    def _get_column_by_expected_name(self, df, expected_column_name):
        if df is None or not expected_column_name:
            return None

        expected = self._normalize_column_name(expected_column_name)
        for column in df.columns:
            if self._normalize_column_name(column) == expected:
                return column
        return None

    def _get_german_value_for_field(self, sku, field_mapping):
        if not sku or self.german_df is None or self.trackvia_df is None or not field_mapping:
            return ""

        trackvia_sku_column = self._get_mapped_column(self.trackvia_df, "trackvia")
        trackvia_german_part_column = self._get_column_by_expected_name(self.trackvia_df, "part_german")
        german_part_number_column = self._get_column_by_expected_name(self.german_df, "item no.")
        german_field_column = self._get_column_by_expected_name(self.german_df, field_mapping.get("german"))

        if trackvia_sku_column is None or trackvia_german_part_column is None or german_part_number_column is None or german_field_column is None:
            return ""

        for _, row in self.trackvia_df.iterrows():
            current_sku = str(row[trackvia_sku_column]).strip() if pd.notna(row[trackvia_sku_column]) else ""
            if current_sku != sku:
                continue

            german_part_number = str(row[trackvia_german_part_column]).strip() if pd.notna(row[trackvia_german_part_column]) else ""
            if not german_part_number:
                return ""

            for _, german_row in self.german_df.iterrows():
                german_row_part_number = str(german_row[german_part_number_column]).strip() if pd.notna(german_row[german_part_number_column]) else ""
                if german_row_part_number != german_part_number:
                    continue

                value = german_row[german_field_column]
                return str(value).strip() if pd.notna(value) else ""

            break

        return ""

    def _build_recommendation_mismatches(self, mismatches):
        enriched_mismatches = []
        for mismatch in mismatches:
            field_mapping = self.field_mapper.get_mapping(mismatch.get("field_key")) if mismatch.get("field_key") else None
            source_of_truth = field_mapping.get("source_of_truth", "") if field_mapping else ""
            german_value = self._get_german_value_for_field(mismatch.get("sku", ""), field_mapping)

            enriched_mismatches.append(
                {
                    **mismatch,
                    "source_of_truth": source_of_truth,
                    "german_value": german_value,
                }
            )

        return enriched_mismatches

    def _build_product_corrections(self, recommended_corrections, four_way_comparisons, missing_from_directus, missing_from_trackvia):
        sku_lookup = {}

        for record in self.records:
            sku = str(record.sku).strip() if record.sku is not None else ""
            if not sku:
                continue
            sku_lookup[sku] = {
                "sku": sku,
                "product_identity_corrections": [],
                "engineering_corrections": [],
                "website_qa_corrections": [],
                "directus_qa_corrections": [],
                "total_corrections": 0,
            }

        for sku in missing_from_directus:
            if sku not in sku_lookup:
                sku_lookup[sku] = {
                    "sku": sku,
                    "product_identity_corrections": [],
                    "engineering_corrections": [],
                    "website_qa_corrections": [],
                    "directus_qa_corrections": [],
                    "total_corrections": 0,
                }
            sku_lookup[sku]["product_identity_corrections"].append(
                {
                    "system": "Directus",
                    "field": "Part Number",
                    "current_value": "",
                    "correct_value": sku,
                    "reason": "SKU missing from Directus",
                }
            )

        for sku in missing_from_trackvia:
            if sku not in sku_lookup:
                sku_lookup[sku] = {
                    "sku": sku,
                    "product_identity_corrections": [],
                    "engineering_corrections": [],
                    "website_qa_corrections": [],
                    "directus_qa_corrections": [],
                    "total_corrections": 0,
                }
            sku_lookup[sku]["product_identity_corrections"].append(
                {
                    "system": "TrackVia",
                    "field": "Part Number",
                    "current_value": "",
                    "correct_value": sku,
                    "reason": "SKU missing from TrackVia",
                }
            )

        for correction in recommended_corrections:
            sku = correction.get("sku", "")
            if not sku:
                continue
            if sku not in sku_lookup:
                sku_lookup[sku] = {
                    "sku": sku,
                    "product_identity_corrections": [],
                    "engineering_corrections": [],
                    "website_qa_corrections": [],
                    "directus_qa_corrections": [],
                    "total_corrections": 0,
                }
            sku_lookup[sku]["engineering_corrections"].append(
                {
                    "system": correction.get("system", ""),
                    "field": correction.get("field", ""),
                    "current_value": correction.get("current_value", ""),
                    "correct_value": correction.get("correct_value", ""),
                    "reason": correction.get("reason", ""),
                }
            )

        for comparison in four_way_comparisons:
            sku = comparison.get("sku", "")
            analysis = comparison.get("analysis", {})
            systems_out = analysis.get("systems_out_of_sync", [])
            if not sku or not systems_out:
                continue

            if sku not in sku_lookup:
                sku_lookup[sku] = {
                    "sku": sku,
                    "product_identity_corrections": [],
                    "engineering_corrections": [],
                    "website_qa_corrections": [],
                    "directus_qa_corrections": [],
                    "total_corrections": 0,
                }

            if "US Catalog" in systems_out:
                sku_lookup[sku]["website_qa_corrections"].append(
                    {
                        "system": "US Catalog",
                        "field": comparison.get("field", ""),
                        "current_value": comparison.get("us_catalog_value", ""),
                        "correct_value": analysis.get("correct_value", ""),
                        "reason": analysis.get("recommendation", ""),
                    }
                )

            if "Directus" in systems_out:
                sku_lookup[sku]["directus_qa_corrections"].append(
                    {
                        "system": "Directus",
                        "field": comparison.get("field", ""),
                        "current_value": comparison.get("directus_value", ""),
                        "correct_value": analysis.get("correct_value", ""),
                        "reason": analysis.get("recommendation", ""),
                    }
                )

        product_corrections = []
        for sku in sorted(sku_lookup.keys()):
            item = sku_lookup[sku]
            item["total_corrections"] = (
                len(item["product_identity_corrections"])
                + len(item["engineering_corrections"])
                + len(item["website_qa_corrections"])
                + len(item["directus_qa_corrections"])
            )
            product_corrections.append(item)

        return product_corrections

    def _prepare_audit_records_for_outlier_analysis(self):
        required_fields = ["sku", "german_part", "german_row", "us_catalog_row", "trackvia_row", "directus_row"]
        complete_records = []

        for record in self.records:
            # Keep the full AuditRecord object while verifying the required attributes are present.
            if all(hasattr(record, field_name) for field_name in required_fields):
                complete_records.append(record)

        return complete_records

    def _find_missing_part_numbers(self, trackvia_df, directus_df):
        trackvia_column = self._get_mapped_column(trackvia_df, "trackvia")
        directus_column = self._get_mapped_column(directus_df, "directus")

        if trackvia_column is None or directus_column is None:
            return [], []

        trackvia_values = self._unique_values(trackvia_df, trackvia_column)
        directus_values = self._unique_values(directus_df, directus_column)

        directus_lookup = set(directus_values)
        missing_from_directus = [value for value in trackvia_values if value not in directus_lookup]

        trackvia_lookup = set(trackvia_values)
        missing_from_trackvia = [value for value in directus_values if value not in trackvia_lookup]

        return missing_from_directus, missing_from_trackvia

    def _build_tier1_field_summary(self, mismatches):
        summary = {}
        trackvia_column = self._get_mapped_column(self.trackvia_df, "trackvia")
        directus_column = self._get_mapped_column(self.directus_df, "directus")

        common_sku_count = 0
        if trackvia_column is not None and directus_column is not None:
            trackvia_skus = set(self._unique_values(self.trackvia_df, trackvia_column))
            directus_skus = set(self._unique_values(self.directus_df, directus_column))
            common_sku_count = len(trackvia_skus & directus_skus)

        for field in self.field_mapper.get_specification_mappings():
            field_key = field.get("key")
            field_display = field.get("display")
            if not field_key or not field_display:
                continue

            field_mismatches = [mismatch for mismatch in mismatches if mismatch.get("field_key") == field_key]
            summary[field_display] = {
                "comparisons": common_sku_count,
                "mismatches": len(field_mismatches),
                "status": "PASS" if len(field_mismatches) == 0 else "FAIL",
            }

        return summary

    def _build_part_number_comparison_html(self):
        trackvia_column = self._get_mapped_column(self.trackvia_df, "trackvia")
        directus_column = self._get_mapped_column(self.directus_df, "directus")

        warnings = []
        if trackvia_column is None:
            available_columns = ", ".join(str(column) for column in self.trackvia_df.columns) if self.trackvia_df is not None else ""
            warnings.append(
                f"TrackVia expected column 'sku' is missing. Available columns: {available_columns or 'none'}"
            )
        if directus_column is None:
            available_columns = ", ".join(str(column) for column in self.directus_df.columns) if self.directus_df is not None else ""
            warnings.append(
                f"Directus expected column 'sku' is missing. Available columns: {available_columns or 'none'}"
            )

        if warnings:
            return (
                '<div class="alert alert-warning mt-4" role="alert">'
                + " ".join(warnings)
                + "</div>"
            )

        missing_from_directus, missing_from_trackvia = self._find_missing_part_numbers(self.trackvia_df, self.directus_df)
        matching_count = len(set(self._unique_values(self.trackvia_df, trackvia_column)) & set(self._unique_values(self.directus_df, directus_column)))
        mismatches = compare_specifications(self.trackvia_df, self.directus_df, self.field_mapper)
        part_number_mapping = build_part_number_mapping(self.trackvia_df)

        missing_from_directus_text = "\n".join(missing_from_directus) if missing_from_directus else "None"
        missing_from_trackvia_text = "\n".join(missing_from_trackvia) if missing_from_trackvia else "None"

        duplicate_german_numbers = []
        duplicate_german_lookup = {}
        duplicate_us_skus = []
        us_sku_lookup = {}

        german_column = None
        sku_column = None
        if self.trackvia_df is not None:
            german_column = next((column for column in self.trackvia_df.columns if self._normalize_column_name(column) == "partgerman"), None)
            sku_column = next((column for column in self.trackvia_df.columns if self._normalize_column_name(column) == "sku"), None)

        if german_column is not None and sku_column is not None:
            for _, row in self.trackvia_df.iterrows():
                german_value = str(row[german_column]).strip() if pd.notna(row[german_column]) else ""
                sku_value = str(row[sku_column]).strip() if pd.notna(row[sku_column]) else ""

                if not german_value:
                    continue

                if german_value in duplicate_german_lookup:
                    if german_value not in duplicate_german_numbers:
                        duplicate_german_numbers.append(german_value)
                else:
                    duplicate_german_lookup[german_value] = sku_value

                us_sku_lookup.setdefault(sku_value, []).append(german_value)

            for sku, german_numbers in us_sku_lookup.items():
                if sku and len(set(german_numbers)) > 1:
                    duplicate_us_skus.append((sku, sorted(set(german_numbers))))

        mapping_rows = "".join(
            f"<tr><td>{german_number}</td><td>{sku}</td></tr>"
            for german_number, sku in part_number_mapping.items()
        )

        tier1_summary = self._build_tier1_field_summary(mismatches)
        tier1_rows = "".join(
            f"<tr><td>{field_name}</td><td>{details['comparisons']}</td><td>{details['mismatches']}</td><td>{details['status']}</td></tr>"
            for field_name, details in tier1_summary.items()
        )
        tier1_table = (
            '<div class="card shadow-sm mt-4">'
            '<div class="card-body">'
            '<h2 class="h5 mb-3">Tier 1 Field Summary</h2>'
            + (f'<div class="table-responsive"><table class="table table-striped align-middle"><thead><tr><th>Field</th><th>Comparisons</th><th>Mismatches</th><th>Status</th></tr></thead><tbody>{tier1_rows}</tbody></table></div>' if tier1_rows else '<p class="mb-0">No tier 1 fields found.</p>')
            + '</div></div>'
        )

        recommendation_payload = self._build_recommendation_mismatches(mismatches)
        correction_rows = "".join(
            f"<tr><td>{recommendation['sku']}</td><td>{recommendation['system']}</td><td>{recommendation['field']}</td><td>{recommendation['current_value']}</td><td>{recommendation['correct_value']}</td><td>{recommendation['reason']}</td><td>{recommendation.get('source_of_truth', '')}</td></tr>"
            for recommendation in build_correction_recommendations(recommendation_payload)[:20]
        )
        correction_table = (
            '<div class="card shadow-sm mt-4">'
            '<div class="card-body">'
            '<h2 class="h5 mb-3">Recommended Corrections</h2>'
            + (f'<div class="table-responsive"><table class="table table-striped align-middle"><thead><tr><th>SKU</th><th>System</th><th>Field</th><th>Current Value</th><th>Correct Value</th><th>Reason</th><th>Source of Truth</th></tr></thead><tbody>{correction_rows}</tbody></table></div>' if correction_rows else '<p class="mb-0">No recommended corrections found.</p>')
            + '</div></div>'
        )

        mismatch_rows = "".join(
            f"<tr><td>{mismatch['sku']}</td><td>{mismatch['field_display']}</td><td>{mismatch['trackvia_value']}</td><td>{mismatch['directus_value']}</td></tr>"
            for mismatch in mismatches
        )
        mismatch_table = (
            '<div class="card shadow-sm mt-4">'
            '<div class="card-body">'
            '<h2 class="h5 mb-3">Specification Mismatches</h2>'
            f'<p class="mb-3"><strong>Total mismatches:</strong> {len(mismatches)}</p>'
            + (f'<div class="table-responsive"><table class="table table-striped align-middle"><thead><tr><th>SKU</th><th>Field</th><th>TrackVia</th><th>Directus</th></tr></thead><tbody>{mismatch_rows}</tbody></table></div>' if mismatch_rows else '<p class="mb-0">No specification mismatches found.</p>')
            + '</div></div>'
        )

        mapping_warning_html = ""
        if duplicate_german_numbers or duplicate_us_skus:
            warning_items = []
            if duplicate_german_numbers:
                warning_items.append(
                    f"Duplicate German part numbers found: {', '.join(duplicate_german_numbers)}"
                )
            if duplicate_us_skus:
                formatted_duplicates = ", ".join(
                    f"{sku} -> {', '.join(numbers)}" for sku, numbers in duplicate_us_skus
                )
                warning_items.append(f"Duplicate US SKUs mapped from different German numbers: {formatted_duplicates}")
            mapping_warning_html = (
                '<div class="alert alert-warning mt-3" role="alert">'
                + "<strong>Part number mapping warnings:</strong> "
                + " | ".join(warning_items)
                + "</div>"
            )

        mapping_table = (
            '<div class="card shadow-sm mt-4">'
            '<div class="card-body">'
            '<h2 class="h5 mb-3">German ↔ US Part Number Mapping</h2>'
            + mapping_warning_html
            + (f'<div class="table-responsive"><table class="table table-striped align-middle"><thead><tr><th>German Part Number</th><th>US Part Number</th></tr></thead><tbody>{mapping_rows}</tbody></table></div>' if mapping_rows else '<p class="mb-0">No part number mappings found.</p>')
            + '</div></div>'
        )

        return (
            '<div class="card shadow-sm mt-4">'
            '<div class="card-body">'
            '<h2 class="h5 mb-3">Part Number Comparison</h2>'
            '<div class="row g-3">'
            '<div class="col-12">'
            '<p class="mb-2"><strong>Matching part numbers:</strong> '
            + str(matching_count)
            + '</p>'
            '<p class="mb-2"><strong>Missing from Directus:</strong> '
            + str(len(missing_from_directus))
            + '</p>'
            '<p class="mb-2"><strong>Missing from TrackVia:</strong> '
            + str(len(missing_from_trackvia))
            + '</p>'
            '</div>'
            '<div class="col-md-6">'
            '<h3 class="h6">Missing from Directus</h3>'
            '<pre class="border rounded p-3 bg-light mb-0">'
            + missing_from_directus_text
            + '</pre></div>'
            '<div class="col-md-6">'
            '<h3 class="h6">Missing from TrackVia</h3>'
            '<pre class="border rounded p-3 bg-light mb-0">'
            + missing_from_trackvia_text
            + '</pre></div></div></div></div>'
            + mapping_table
            + tier1_table
            + correction_table
            + mismatch_table
        )

    def run(self):
        comparisons = FourWayComparison().compare(self.records)
        self.four_way_comparisons = comparisons

        report_filename = ""

        if self.trackvia_df is not None and self.directus_df is not None:
            missing_from_directus, missing_from_trackvia = self._find_missing_part_numbers(self.trackvia_df, self.directus_df)
            mismatches = compare_specifications(self.trackvia_df, self.directus_df, self.field_mapper)
            part_number_mapping = build_part_number_mapping(self.trackvia_df)
            trackvia_column = self._get_mapped_column(self.trackvia_df, "trackvia")
            directus_column = self._get_mapped_column(self.directus_df, "directus")
            tier1_summary = self._build_tier1_field_summary(mismatches)

            summary_data = {
                "audit_type": self.audit_type,
                "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "matching_count": len(
                    set(self._unique_values(self.trackvia_df, trackvia_column))
                    & set(self._unique_values(self.directus_df, directus_column))
                ),
                "missing_from_directus_count": len(missing_from_directus),
                "missing_from_trackvia_count": len(missing_from_trackvia),
                "total_mismatches": len(mismatches),
                "missing_from_directus": missing_from_directus,
                "missing_from_trackvia": missing_from_trackvia,
                "mismatches": mismatches,
                "part_number_mapping": part_number_mapping,
                "tier1_summary": tier1_summary,
                "recommended_corrections": [],
                "product_corrections": [],
            }
            summary_data["four_way_comparisons"] = self.four_way_comparisons
            if summary_data["four_way_comparisons"]:
                print("================ FIRST SUMMARY COMPARISON ================")
                print()
                print(summary_data["four_way_comparisons"][0])
                print()
                print("==========================================================")
            print("EXCEL MODULE:")
            print(excel_report_module.__file__)
            report_filename = excel_report_module.write_audit_report(summary_data)

        return {
            "comparison_html": "",
            "report_filename": report_filename,
            "four_way_comparisons": self.four_way_comparisons,
        }
